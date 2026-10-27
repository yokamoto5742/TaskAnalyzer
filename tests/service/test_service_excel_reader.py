

import configparser
import tempfile
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from service.service_excel_reader import ExcelTaskReader, NoDataError, parse_sheet_date


@pytest.fixture
def mock_config():
    config = configparser.ConfigParser()
    config['Analysis'] = {
        'start_row': '5',
        'end_row': '15',
        'daily_task_start_row': '20',
        'daily_task_end_row': '25',
        'communication_start_row': '30',
        'communication_end_row': '35'
    }
    return config


@pytest.fixture
def mock_workbook():
    # テスト用のExcelファイルを作成
    wb = Workbook()
    
    # デフォルトシートを削除
    wb.remove(wb.worksheets[0])
    
    # 3日分のシートを作成
    for day in range(1, 4):
        date_str = f"2024年1月{day}日"
        sheet = wb.create_sheet(title=f'シート{day}')
        
        # 日付を設定
        sheet['A1'] = date_str
        
        # クラーク業務データ
        if day == 1:
            sheet['B5'] = 'クラーク業務A'
            sheet['C5'] = 30
            sheet['B6'] = 'クラーク業務B'
            sheet['C6'] = 45
        elif day == 2:
            sheet['B5'] = 'クラーク業務A'
            sheet['C5'] = 25
            sheet['B6'] = '会議'
            sheet['C6'] = 60
        else:
            sheet['B5'] = '資料作成'
            sheet['C5'] = 90
        
        # デイリータスクデータ
        if day in [1, 2]:
            sheet['B20'] = '毎日タスクA'
            sheet['C20'] = 10
        if day in [1, 3]:
            sheet['B21'] = '毎日タスクB'
            sheet['C21'] = 15
        
        # コミュニケーションデータ
        if day == 1:
            sheet['B30'] = '打合せ(田中)'
            sheet['C30'] = 30
            sheet['B31'] = 'レビュー(佐藤)'
            sheet['C31'] = 45
        elif day == 2:
            sheet['B30'] = '打合せ(田中)'
            sheet['C30'] = 25
            sheet['B31'] = '相談(鈴木)'
            sheet['C31'] = 15
        else:
            sheet['B30'] = 'レビュー(佐藤)'
            sheet['C30'] = 30
    
    # 特殊なシートを追加
    wb.create_sheet(title='シート一覧')
    
    # 一時ファイルに保存
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
        wb.save(tmp.name)
        return tmp.name


@pytest.fixture
def first_sheet(mock_workbook):
    return load_workbook(filename=mock_workbook)['シート1']


class TestExcelTaskReader:
    def test_extract_cell_data(self, mock_config):
        reader = ExcelTaskReader(mock_config)
        sheet = Workbook().active
        assert sheet is not None
        sheet['B1'], sheet['C1'] = 'テストタスク 補足', 30
        # 無効なデータ
        sheet['B2'], sheet['C2'] = 'テストタスク', '*'
        sheet['B3'], sheet['C3'] = None, 30
        sheet['B4'], sheet['C4'] = 'テストタスク', 'abc'  # 数値に変換できない
        sheet['B5'], sheet['C5'] = '(田中)', 30  # 名前を除くと内容が空になる

        sheet_date = date(2024, 1, 1)
        result = reader.extract_cell_data(sheet, 1, sheet_date)

        assert result == {'date': sheet_date, 'content': 'テストタスク', 'minutes': 30}
        for row in (2, 3, 4):
            assert reader.extract_cell_data(sheet, row, sheet_date) is None
        assert reader.extract_cell_data(sheet, 5, sheet_date, remove_name=True) is None

    def test_load_excel_task_data(self, mock_config, first_sheet):
        tasks, daily_tasks, comm_tasks = ExcelTaskReader(mock_config).load_excel_task_data(
            first_sheet, date(2024, 1, 1))

        assert len(tasks) == 2
        assert tasks[0]['content'] == 'クラーク業務A'
        assert tasks[0]['minutes'] == 30

        assert len(daily_tasks) == 2
        assert daily_tasks[0]['content'] == '毎日タスクA'

        assert len(comm_tasks) == 2
        assert comm_tasks[0]['name'] == '田中'
        assert comm_tasks[0]['content'] == '打合せ'
        assert comm_tasks[0]['minutes'] == 30

    def test_load_excel_task_data_communication_formats(self, mock_config):
        """空白区切りの名前も読み取り、名前のない行はコミュニケーションに含めないこと"""
        sheet = Workbook().active
        assert sheet is not None
        sheet['B30'], sheet['C30'] = '打合せ 補足 (田中)', 20
        sheet['B31'], sheet['C31'] = '雑談', 10

        _, _, comm_tasks = ExcelTaskReader(mock_config).load_excel_task_data(sheet, date(2024, 1, 1))

        assert [(task['content'], task['name']) for task in comm_tasks] == [('打合せ', '田中')]

    def test_load_excel_sheet_all_items(self, mock_config, first_sheet):
        all_items = ExcelTaskReader(mock_config).load_excel_sheet_all_items(first_sheet, date(2024, 1, 1))

        assert len(all_items) == 4  # クラーク業務2つ + デイリータスク2つ

    def test_load_excel_sheet_all_items_removes_name(self, mock_config, first_sheet):
        """全項目に含まれるコミュニケーション行は、空白がなくても名前を除いて集計すること"""
        mock_config['Analysis']['daily_task_end_row'] = '35'

        all_items = ExcelTaskReader(mock_config).load_excel_sheet_all_items(first_sheet, date(2024, 1, 1))

        contents = [item['content'] for item in all_items]
        assert '打合せ' in contents
        assert 'レビュー' in contents
        assert not any('(' in content for content in contents)

    def test_read_workbook(self, mock_config, mock_workbook):
        workbook_data = ExcelTaskReader(mock_config).read_workbook(
            mock_workbook, date(2024, 1, 1), date(2024, 1, 3))

        assert len(workbook_data.tasks) == 5
        assert len(workbook_data.daily_tasks) == 4
        assert len(workbook_data.communication_tasks) == 5
        assert len(workbook_data.all_items) == 9
        assert workbook_data.first_date == date(2024, 1, 1)
        assert workbook_data.last_date == date(2024, 1, 3)

    def test_read_workbook_date_filter(self, mock_config, mock_workbook):
        workbook_data = ExcelTaskReader(mock_config).read_workbook(
            mock_workbook, date(2024, 1, 2), date(2024, 1, 2))

        # 1日分のデータのみ
        assert len(workbook_data.tasks) == 2
        assert workbook_data.first_date == date(2024, 1, 2)
        assert workbook_data.last_date == date(2024, 1, 2)

    def test_read_workbook_no_data(self, mock_config, mock_workbook):
        with pytest.raises(NoDataError):
            ExcelTaskReader(mock_config).read_workbook(mock_workbook, date(2025, 1, 1), date(2025, 1, 31))


@pytest.mark.parametrize("value, expected", [
    ('2024年1月2日', date(2024, 1, 2)),
    (datetime(2024, 1, 2, 9, 30), date(2024, 1, 2)),
    ('目次', None),
    (None, None),
])
def test_parse_sheet_date(value, expected):
    assert parse_sheet_date(value) == expected
