import re
from datetime import datetime
from openpyxl import load_workbook

NAME_PATTERN = re.compile(r'\((.*?)\)')


class ExcelTaskReader:
    def __init__(self, config):
        self.config = config

    @staticmethod
    def extract_cell_data(sheet, row, date, remove_name=False):
        content = str(sheet[f'B{row}'].value or '')
        time = sheet[f'C{row}'].value

        if not content or not time:
            return None

        try:
            minutes = float(time)
        except (ValueError, TypeError):
            return None

        if remove_name:
            content = NAME_PATTERN.sub('', content)
        words = content.split()
        if not words:
            return None

        return {'date': date, 'content': words[0], 'minutes': minutes}

    def _rows(self, start_key, end_key):
        return range(self.config.getint('Analysis', start_key),
                     self.config.getint('Analysis', end_key) + 1)

    def _read_tasks(self, sheet, rows, date):
        tasks = (self.extract_cell_data(sheet, row, date) for row in rows)
        return [task for task in tasks if task]

    def _read_communications(self, sheet, date):
        communications = []
        for row in self._rows('communication_start_row', 'communication_end_row'):
            # 「内容(名前)」形式で名前があるものだけを対象にする
            name_match = NAME_PATTERN.search(str(sheet[f'B{row}'].value))
            task = self.extract_cell_data(sheet, row, date, remove_name=True)
            if name_match and task:
                communications.append({**task, 'name': name_match.group(1)})
        return communications

    def load_excel_task_data(self, wb, sheet_name, date):
        sheet = wb[sheet_name]
        tasks = self._read_tasks(sheet, self._rows('start_row', 'end_row'), date)
        daily_tasks = self._read_tasks(sheet, self._rows('daily_task_start_row', 'daily_task_end_row'), date)
        return tasks, daily_tasks, self._read_communications(sheet, date)

    def load_excel_sheet_all_items(self, wb, sheet_name, date):
        sheet = wb[sheet_name]
        comm_rows = self._rows('communication_start_row', 'communication_end_row')
        items = (self.extract_cell_data(sheet, row, date, remove_name=row in comm_rows)
                 for row in self._rows('start_row', 'daily_task_end_row'))
        return [item for item in items if item]

    def read_workbook(self, file_path, start_date, end_date):
        wb = load_workbook(filename=file_path)
        all_tasks = []
        all_daily_tasks = []
        all_communication_tasks = []
        all_items = []
        dates = []

        for sheet_name in wb.sheetnames:
            if sheet_name == 'シート一覧':
                continue

            sheet = wb[sheet_name]
            date_cell = sheet['A1'].value

            try:
                if isinstance(date_cell, datetime):
                    sheet_date = date_cell
                else:
                    sheet_date = datetime.strptime(str(date_cell), '%Y年%m月%d日')

                if start_date <= sheet_date <= end_date:
                    tasks, daily_tasks, comm_tasks = self.load_excel_task_data(wb, sheet_name, sheet_date)
                    sheet_items = self.load_excel_sheet_all_items(wb, sheet_name, sheet_date)

                    all_tasks.extend(tasks)
                    all_daily_tasks.extend(daily_tasks)
                    all_communication_tasks.extend(comm_tasks)
                    all_items.extend(sheet_items)
                    dates.append(sheet_date)

            except (ValueError, TypeError) as e:
                print(f"シート {sheet_name} の日付の解析でエラー: {e}")
                continue

        if not dates:
            raise ValueError("指定された期間内のデータがありません")

        actual_start_date = min(dates).strftime("%Y%m%d")
        actual_end_date = max(dates).strftime("%Y%m%d")

        return all_tasks, all_daily_tasks, all_communication_tasks, all_items, actual_start_date, actual_end_date
